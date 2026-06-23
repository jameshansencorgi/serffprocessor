from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path


@dataclass
class PageText:
    page_number: int
    text: str
    ocr_used: bool = False


@dataclass
class ParsedDocument:
    pages: list[PageText]
    parser: str
    failures: list[str]


def parse_document(path: Path) -> ParsedDocument:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _parse_pdf(path)
    if suffix in {".txt", ".md", ".csv"}:
        return _parse_plain_text(path)
    if suffix == ".docx":
        return _parse_docx(path)
    if suffix in {".xlsx", ".xlsm"}:
        return _parse_excel(path)
    return ParsedDocument([], "unsupported", [f"unsupported extension: {suffix}"])


def _parse_plain_text(path: Path) -> ParsedDocument:
    text = path.read_text(errors="ignore")
    return ParsedDocument([PageText(1, text)], "plain_text", [])


def _parse_pdf(path: Path) -> ParsedDocument:
    failures: list[str] = []
    try:
        import pdfplumber

        pages: list[PageText] = []
        with pdfplumber.open(path) as pdf:
            for idx, page in enumerate(pdf.pages, start=1):
                text = page.extract_text() or ""
                pages.append(PageText(idx, text))
        if any(page.text.strip() for page in pages):
            return ParsedDocument(pages, "pdfplumber", failures)
        failures.append("pdfplumber extracted no text; OCR fallback is not configured in MVP")
        return ParsedDocument(pages, "pdfplumber", failures)
    except Exception as exc:
        failures.append(f"pdf parse failed: {exc}")
    return ParsedDocument([], "pdfplumber", failures)


def _parse_docx(path: Path) -> ParsedDocument:
    try:
        from docx import Document

        doc = Document(path)
        parts: list[str] = []
        parts.extend(p.text for p in doc.paragraphs if p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return ParsedDocument([PageText(1, "\n".join(parts))], "python-docx", [])
    except Exception as exc:
        return ParsedDocument([], "python-docx", [f"docx parse failed: {exc}"])


def _parse_excel(path: Path) -> ParsedDocument:
    try:
        import openpyxl

        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        pages: list[PageText] = []
        for idx, sheet in enumerate(wb.worksheets, start=1):
            lines: list[str] = [f"Sheet: {sheet.title}"]
            for row in sheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                if any(value.strip() for value in values):
                    lines.append(" | ".join(values))
            pages.append(PageText(idx, "\n".join(lines)))
        return ParsedDocument(pages, "openpyxl", [])
    except Exception as exc:
        return ParsedDocument([], "openpyxl", [f"excel parse failed: {exc}"])

